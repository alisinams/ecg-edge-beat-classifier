"""Render `Reduced Version 1.md` as a Word document with its figures in place.

The markdown already carries the tables, the figure captions, the display
equations and the reference list, because `29_build_reduced.py` spliced them in
verbatim and verified them. This step only has to do three things: convert the
markdown with the same converter the full manuscript uses, so equations become
real Word equation objects rather than LaTeX source; drop each figure image
above its caption; and set the reference entries in a smaller face with a
hanging indent so the list reads as a bibliography and not as body text.
"""
from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
CODE = Path(__file__).resolve().parent
FIGDIR = ROOT / "Folder Figures"
MD = ROOT / "Reduced Version 1.md"
DOCX = ROOT / "Reduced Version 1.docx"


def _load(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


builder = _load("builder", CODE / "17_build_docx.py")

FIGURES = [
    ("Figure 1.", "Figure_1_pipeline", 6.0),
    ("Figure 2.", "Figure_2_confusion_matrices", 6.0),
    ("Figure 3.", "Figure_3_precision_recall", 6.0),
    ("Figure 4.", "Figure_4_reliability", 4.4),
    ("Figure 5.", "Figure_5_coverage", 4.4),
    ("Figure 6.", "Figure_6_noise", 5.4),
    ("Figure 7.", "Figure_7_size_accuracy", 6.0),
]

REF_LINE = re.compile(r"\[#\d+\]\s*$")


def place_figures(doc) -> int:
    placed = 0
    for prefix, stem, width_in in FIGURES:
        png = FIGDIR / f"{stem}.png"
        if not png.exists():
            print(f"  MISSING {png.name}")
            continue
        target = next((p for p in doc.paragraphs
                       if p.text.strip().startswith(prefix)
                       and len(p.text.strip()) > 60), None)
        if target is None:
            print(f"  no caption paragraph for {prefix}")
            continue
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.space_after = Pt(4)
        par.add_run().add_picture(str(png), width=Inches(width_in))
        target._p.addprevious(par._p)
        placed += 1
    return placed


def style_references(doc) -> int:
    n = 0
    for p in doc.paragraphs:
        if not REF_LINE.search(p.text.strip()):
            continue
        pf = p.paragraph_format
        pf.left_indent = Inches(0.3)
        pf.first_line_indent = Inches(-0.3)
        pf.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(9)
        n += 1
    return n


def main() -> int:
    if not MD.exists():
        print(f"  {MD.name} not built yet; run 29_build_reduced.py first")
        return 1

    builder._init_math()
    builder.convert(MD, DOCX)
    print(f"  converted {MD.name} -> {DOCX.name}")
    print(f"  equations converted: {builder.CONVERTED}, "
          f"fell back to source: {builder.FALLBACKS}")

    doc = Document(DOCX)
    figs = place_figures(doc)
    refs = style_references(doc)
    doc.save(DOCX)

    doc = Document(DOCX)
    print(f"  figures placed:   {figs} of {len(FIGURES)}")
    print(f"  tables:           {len(doc.tables)}")
    print(f"  reference entries {refs}")
    print(f"  images in file:   {len(doc.inline_shapes)}")
    print(f"  size:             {DOCX.stat().st_size / 1024:.0f} kB")
    ok = figs == len(FIGURES) and len(doc.tables) == 9 and refs == 92
    print("REDUCED_DOCX_COMPLETE" if ok
          else "REDUCED_DOCX_INCOMPLETE, see the counts above")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
