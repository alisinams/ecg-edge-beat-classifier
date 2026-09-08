"""Place the rendered figures into the Word manuscript and supplement.

The manuscript is read from `08_Manuscript_v3.docx` and written to
`08_Manuscript_v4.docx`, so the version that arrived without figures is left
exactly as it was. The supplement is written in place, with its figure-free
state kept beside it as `08_Supplementary.nofigures.docx`.

Each figure is inserted as its own centred paragraph immediately above the
caption that already sits in the document, so caption numbering, order and
wording are untouched. The caption paragraph is the anchor and is matched on
its opening words, not on an index, so the script stays correct as paragraphs
shift during the run.

Widths are set in inches at insertion time. The text column of these documents
is 6.0 inches. A 180 mm figure is placed at the full column width, which is a
15 per cent reduction from its drawn size; the single-column figures are placed
larger than their drawn width so that their type reads comfortably on screen
and on paper. The PNGs are 600 dpi, so enlarging costs no sharpness. The PDFs
in `Folder Figures` remain the masters at their exact journal widths.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import importlib.util

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

_spec = importlib.util.spec_from_file_location(
    "refs", Path(__file__).resolve().parent / "26_refs.py")
refs = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(refs)

ROOT = Path(__file__).resolve().parent.parent
FIGDIR = ROOT / "Folder Figures"

# (caption opening, figure file stem, placed width in inches)
MAIN = [
    ("Figure 1.", "Figure_1_pipeline", 6.0),
    ("Figure 2.", "Figure_2_confusion_matrices", 6.0),
    ("Figure 3.", "Figure_3_precision_recall", 6.0),
    ("Figure 4.", "Figure_4_reliability", 4.4),
    ("Figure 5.", "Figure_5_coverage", 4.4),
    ("Figure 6.", "Figure_6_noise", 5.4),
    ("Figure 7.", "Figure_7_size_accuracy", 6.0),
]

SUPP = [
    ("Figure S4.1.", "Figure_S4-1_optimiser_convergence", 6.0),
    ("Figure S4.2.", "Figure_S4-2_roc", 6.0),
    ("Figure S4.3.", "Figure_S4-3_external_confusion", 5.4),
    ("Figure S4.4.", "Figure_S4-4_protocol_slopegraph", 5.6),
]

# Text corrections that follow from drawing Figure 7 from `fig7_frontier.csv`.
# The counts in the prose did not match the file the figure is built from, and
# the shaded region turned out to hold one anchor point rather than none. The
# same edits are in `05_Manuscript.md`; they are repeated here because the Word
# manuscript is edited directly rather than rebuilt from the markdown.
MANUSCRIPT_PATCHES = [
    ("Of the studies listed, 3 partition by patient, 4 state a model size in "
     "bytes, and 5 state an energy figure.",
     "Of the seven edge-deployment studies, 3 partition by patient, 4 state a "
     "model size in bytes, and 3 state an energy figure."),

    ("The point of the figure is the empty region of the plane,",
     "The point of the figure is how thinly the plane is populated,"),

    ("Marker shape encodes the evaluation protocol: filled markers are studies "
     "that partitioned by patient, hollow markers are studies that did not or "
     "did not say, and the distinction matters more than the vertical position "
     "of any point. Marker size encodes reported energy per inference where a "
     "study reports one, and points with no energy figure are drawn at a fixed "
     "minimum size with an open centre, so the reader can count how many "
     "anchors omit the quantity.",
     "Marker fill encodes the evaluation protocol: filled markers are the 3 "
     "studies that partitioned by patient, hollow markers are the 4 that did "
     "not or did not say, and the distinction matters more than the vertical "
     "position of any point. Marker area encodes reported energy per inference "
     "on a logarithmic scale where a study reports one, and the 4 anchors that "
     "report no energy are drawn at a fixed minimum size with a cross through "
     "the centre, so the reader can count them."),

    ("The upper-left region, small models validated between patients with a "
     "full deployment profile, holds almost no published points, and that "
     "emptiness is the gap this study was designed to occupy.",
     "The shaded region, models under 20 kB reported above 95 per cent "
     "accuracy, holds exactly one anchor point, and the qualifications on that "
     "point are the reason it is drawn with a dotted outline: its headline "
     "accuracy is a three-class variant whose five-class result on the same "
     "data is 54.90 per cent average F1, and it reports no energy "
     "{Farag, 2023 #36}. No study in this set, this one included, occupies "
     "that region with a five-class AAMI result obtained between patients and "
     "a complete deployment profile attached. That absence, and not the "
     "position of any marker, is the gap this study was designed to occupy."),

    ("Two of the seven do not report a model size in bytes at all, two report "
     "no energy figure, and three do not partition by patient.",
     "Three of the seven do not report a model size in bytes at all, four "
     "report no energy per inference, and four do not partition by patient or "
     "do not say that they do."),

    # Two factual corrections raised by the citation audit and confirmed
    # against the manuscript's own Table 1. The size span was overstated by
    # two orders of magnitude, and a two-class detector was described as a
    # five-class study.
    ("Reported model sizes in this group span more than three orders of "
     "magnitude, so a size claim carries no information without a named target "
     "device and toolchain.",
     "The four anchors that state a size in bytes span 15 to 64 kB, a factor "
     "of four, and the other three state no byte figure at all, reporting "
     "lookup-table counts or die area instead, so a size claim in this "
     "literature carries no information without a named target device and "
     "toolchain."),

    ("The studies in that group report five-class results on the same "
     "partition of the same database, and the classifier reported here sits "
     "below them {Zhou, 2024 #27; Midani, 2024 #28; Chen, 2024 #29; "
     "Mommen, 2026 #42}. Three differences account for part of the gap and "
     "none of them excuses it. Two of those studies drop the Q superclass or "
     "merge classes, which raises a macro average that this paper computes "
     "over five classes including one with 7 test beats.",
     "Two of those studies report five-class results on the same partition of "
     "the same database and the classifier reported here sits below both "
     "{Zhou, 2024 #27; Midani, 2024 #28}. The other two are not five-class "
     "studies and the comparison has to be stated differently: one drops the Q "
     "superclass and reports four classes {Mommen, 2026 #42}, and one is a "
     "two-class premature-ventricular-contraction detector, so no "
     "macro-averaged figure of ours can be set against it at all "
     "{Chen, 2024 #29}. Three differences account for part of the gap and none "
     "of them excuses it. Dropping or merging classes raises a macro average "
     "that this paper computes over five classes including one with 7 test "
     "beats."),
]


def add_bibliography(doc, before_heading: str = "12. Supplementary materials") -> int:
    """Write the formatted reference list into the References section.

    The list is placed at the end of section 11, which is where a journal
    expects it and which is the end of the paper proper. Entries are set with a
    hanging indent so that the author name of each one starts at the margin.
    Each entry keeps its EndNote record number in brackets, so a reader can
    match {Author, Year #N} to its line, and re-attaching the library in Word
    still regenerates the list in the journal's own style.
    """
    records = refs.parse_ris(refs.RIS)
    entries = refs.bibliography(records)

    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith(before_heading):
            anchor = p
            break

    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(8)
    r = lead.add_run(
        f"The {len(entries)} records of the accompanying library follow, "
        f"alphabetical by first author. The bracketed number ending each entry "
        f"is the EndNote record number, so an in-text marker of the form "
        f"{{Author, Year #N}} maps to its entry without ambiguity.")
    r.italic = True
    r.font.size = Pt(9)
    if anchor is not None:
        anchor._p.addprevious(lead._p)

    for entry in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.3)
        pf.first_line_indent = Inches(-0.3)     # hanging indent
        pf.space_after = Pt(4)
        run = p.add_run(entry)
        run.font.size = Pt(9)
        if anchor is not None:
            anchor._p.addprevious(p._p)

    return len(entries)


def replace_in_paragraph(par, old: str, new: str) -> bool:
    """Swap one span of text, keeping the run formatting around it intact.

    The span usually sits inside a single run, but a citation or a bold word
    can split it. Runs are walked over their original lengths, the replacement
    lands in the run where the match starts, and any tail runs keep only the
    part beyond the match.
    """
    texts = [r.text for r in par.runs]
    i = "".join(texts).find(old)
    if i < 0:
        return False
    j = i + len(old)

    pos = 0
    for run, original in zip(par.runs, texts):
        start, end = pos, pos + len(original)
        pos = end
        if end <= i or start >= j:
            continue
        head = original[:i - start] if start <= i else ""
        tail = original[j - start:] if end > j else ""
        run.text = (head + new + tail) if start <= i < end else tail
    return True


def apply_patches(doc, patches) -> int:
    applied = 0
    for old, new in patches:
        hit = False
        for par in doc.paragraphs:
            if replace_in_paragraph(par, old, new):
                hit = True
                break
        if hit:
            applied += 1
            print(f"  text fixed: \"{old[:56]}...\"")
        else:
            print(f"  TEXT NOT FOUND, left alone: \"{old[:56]}...\"")
    return applied


def already_has_images(doc) -> int:
    return len(doc.inline_shapes)


def insert(src: Path, dst: Path, jobs, label: str, patches=(),
           bibliography: bool = False) -> int:
    doc = Document(src)

    n_existing = already_has_images(doc)
    if n_existing:
        print(f"  {src.name}: {n_existing} images already present; "
              f"expected a figure-free source")
        return -1

    if patches:
        apply_patches(doc, patches)

    if bibliography:
        n = add_bibliography(doc)
        print(f"  reference list added: {n} entries at the end of section 11")

    placed = 0
    for prefix, stem, width_in in jobs:
        png = FIGDIR / f"{stem}.png"
        if not png.exists():
            print(f"  MISSING {png.name}")
            continue

        target = None
        for p in doc.paragraphs:
            t = p.text.strip()
            # the caption, not the sentence in the body that cites the figure
            if t.startswith(prefix) and len(t) > 60:
                target = p
                break
        if target is None:
            print(f"  no caption paragraph found for {prefix}")
            continue

        pic_par = doc.add_paragraph()
        pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_par.paragraph_format.space_before = Pt(10)
        pic_par.paragraph_format.space_after = Pt(4)
        pic_par.add_run().add_picture(str(png), width=Inches(width_in))

        target._p.addprevious(pic_par._p)      # move it above its caption
        placed += 1
        print(f"  {label}: {stem}.png at {width_in}\" above \"{prefix}\"")

    try:
        doc.save(dst)
    except PermissionError:
        print(f"  {dst.name} is open in Word and could not be written; "
              f"close it and run this again")
        return -1
    print(f"  {src.name} -> {dst.name}, {placed} figures")
    return placed


# (source, destination, jobs, label). The manuscript gains a version number
# rather than being overwritten; the supplement keeps a figure-free copy.
JOBS = [
    (ROOT / "08_Manuscript_v3.docx", ROOT / "08_Manuscript_v4.docx",
     MAIN, "manuscript", MANUSCRIPT_PATCHES, True),
    (ROOT / "08_Supplementary.nofigures.docx", ROOT / "08_Supplementary.docx",
     SUPP, "supplement", (), False),
]


def main() -> int:
    total = 0
    for src, dst, jobs, label, patches, bib in JOBS:
        if not src.exists():
            print(f"  skip {dst.name}, source {src.name} not found")
            continue
        n = insert(src, dst, jobs, label, patches, bib)
        if n < 0:
            return 1
        total += n

    for path in sorted(ROOT.glob("08_*.docx")):
        d = Document(path)
        print(f"  {path.name}: {len(d.inline_shapes)} images, "
              f"{path.stat().st_size / 1024:.0f} kB")
    print(f"INSERT_COMPLETE {total} figures placed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
