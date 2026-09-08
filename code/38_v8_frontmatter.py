"""Set the new title and the author block on 12_Manuscript_v8.docx.

The BSPC package carries its own title page; this puts the same title, authors,
affiliation and corresponding author onto the full-length v8 manuscript so the
two do not drift apart.
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "12_Manuscript_v8.docx"

TITLE = ("A Deep Learning ECG Arrhythmia Classifier for Microcontroller-Class "
         "Wearables: Inter-Patient Evaluation, External Validation and Embedded "
         "Deployment Analysis")

AFF = "Computer Department, Islamic Azad University, Zanjan Branch, Zanjan, Iran"

AUTHORS = (
    "Authors. Mahdi Bazargani a,* and Alisina Mousavi a. "
    "a " + AFF + ". "
    "Mahdi Bazargani ORCID [[AUTHOR ACTION: ORCID iD; no public record was "
    "found at ORCID, Crossref, DataCite or Europe PMC under this name]]; "
    "Alisina Mousavi ORCID https://orcid.org/0009-0003-8480-7134."
)

CORRESPONDING = (
    "Corresponding author. Mahdi Bazargani, " + AFF + ", Mbzir@iauz.ac.ir. "
    "[[AUTHOR ACTION: full postal address, including street and postcode, and a "
    "telephone number.]] Co-author email: alisina.mousavi@iau.ir."
)

RUNNING_HEAD = ("Running head. Deep learning ECG classifier for "
                "microcontroller-class wearables")


def set_text(par: Paragraph, text: str) -> None:
    runs = par.runs
    if not runs:
        par.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def main() -> int:
    doc = Document(DOC)
    paras = doc.paragraphs

    def replace_starting(prefix: str, text: str) -> None:
        hits = [p for p in paras if p.text.strip().startswith(prefix)]
        assert len(hits) == 1, f"{prefix!r} matched {len(hits)} paragraphs"
        set_text(hits[0], text)
        print(f"set: {prefix}")

    # the title is the paragraph immediately after the "1. Title" heading
    heading = next(i for i, p in enumerate(paras)
                   if p.text.strip() == "1. Title")
    old_title = paras[heading + 1]
    assert old_title.text.startswith("A distilled lightweight ECG"), \
        f"unexpected title paragraph: {old_title.text[:60]!r}"
    set_text(old_title, TITLE)
    print("set: title")

    replace_starting("Authors.", AUTHORS)
    replace_starting("Corresponding author.", CORRESPONDING)
    replace_starting("Running head.", RUNNING_HEAD)

    doc.save(DOC)
    text = "\n".join(p.text for p in Document(DOC).paragraphs)
    assert "—" not in text and "–" not in text, "dash character present"
    print(f"\nwrote {DOC.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
