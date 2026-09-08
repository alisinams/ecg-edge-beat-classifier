"""Build the Biomedical Signal Processing and Control submission package.

Reads `Biomedical Signal Processing and Control/manuscript_source.txt` and emits
the manuscript .docx in BSPC house style: numbered sections, numbered [n]
citations ordered by first appearance, a reference list in that order, and the
four display equations lifted as real Word equations from 12_Manuscript_v8.docx
so no maths is degraded to an image or to ASCII.

Source markup
  @TITLE / @HU <heading>      title; unnumbered heading
  @H1 <t> / @H2 <t>           auto-numbered headings (1, 1.1, ...)
  @P <text>                   paragraph; {#12} or {#12,#8} are citations
  @EQ <text>                  display equation; consumes the next v8 equation
  @TABLECAP <t> / @TABLE      table caption then pipe-delimited rows
  @FIGCAP <t>                 figure caption, collected after the references
  <i>x</i>                    italic run inside a paragraph
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "Biomedical Signal Processing and Control"
SOURCE = OUT_DIR / "manuscript_source.txt"
REF_LIST = ROOT / "13_Reference_List.md"
V8 = ROOT / "12_Manuscript_v8.docx"
DST = OUT_DIR / "02_Manuscript.docx"

# Body-element indices of the four display equations in 12_Manuscript_v8.docx,
# in the order this manuscript uses them: polynomial score, sigmoid, the
# distillation loss, and energy per inference.
V8_EQUATION_INDICES = [75, 76, 86, 105]

M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


# ----------------------------------------------------------------- parsing
def parse_source(text: str) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("@TABLE\n") or line.strip() == "@TABLE":
            rows = []
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith("@"):
                rows.append([c.strip() for c in lines[i].split("|")])
                i += 1
            blocks.append(("TABLE", rows))
            continue
        m = re.match(r"@(\w+)\s*(.*)$", line)
        if not m:
            raise SystemExit(f"unparsed source line {i + 1}: {line[:60]!r}")
        tag, rest = m.group(1), m.group(2)
        if not rest and tag != "EQ":      # payload is on the following line
            i += 1
            rest = lines[i].strip() if i < len(lines) else ""
        blocks.append((tag, rest))
        i += 1
    return blocks


# ------------------------------------------------------------- numbering
def number_headings(blocks):
    h1 = 0
    h2 = 0
    out = []
    for tag, payload in blocks:
        if tag == "H1":
            h1 += 1
            h2 = 0
            out.append((tag, f"{h1}. {payload}"))
        elif tag == "H2":
            h2 += 1
            out.append((tag, f"{h1}.{h2} {payload}"))
        else:
            out.append((tag, payload))
    return out


CITE = re.compile(r"\{(#\d+(?:\s*,\s*#\d+)*)\}")


def collect_citation_order(blocks) -> dict[int, int]:
    order: dict[int, int] = {}
    def scan(t: str):
        for grp in CITE.findall(t):
            for k in re.findall(r"#(\d+)", grp):
                k = int(k)
                if k not in order:
                    order[k] = len(order) + 1
    for tag, payload in blocks:
        if tag == "TABLE":
            for row in payload:
                for cell in row:
                    scan(cell)
        elif isinstance(payload, str):
            scan(payload)
    return order


def apply_citations(text: str, order: dict[int, int]) -> str:
    def sub(m):
        keys = [int(k) for k in re.findall(r"#(\d+)", m.group(1))]
        nums = sorted(order[k] for k in keys)
        return "[" + ",".join(str(n) for n in nums) + "]"
    return CITE.sub(sub, text)


# ------------------------------------------------------------- references
def load_references() -> dict[int, str]:
    refs: dict[int, str] = {}
    for line in REF_LIST.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^(.*?)\s+\[#(\d+)\]\s*$", line.strip())
        if m:
            refs[int(m.group(2))] = m.group(1).strip()
    return refs


# ------------------------------------------------------------------ docx
def add_par(doc, text: str, style: str | None = None, size: int | None = None,
            bold: bool = False, italic: bool = False, align=None):
    p = doc.add_paragraph(style=style)
    for chunk in re.split(r"(<i>.*?</i>)", text):
        if not chunk:
            continue
        if chunk.startswith("<i>"):
            r = p.add_run(chunk[3:-4])
            r.italic = True
        else:
            r = p.add_run(chunk)
            r.italic = italic
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p


def main() -> int:
    blocks = number_headings(parse_source(SOURCE.read_text(encoding="utf-8")))
    order = collect_citation_order(blocks)
    refs = load_references()
    missing = [k for k in order if k not in refs]
    assert not missing, f"cited keys with no reference entry: {missing}"

    v8 = Document(V8)
    v8_bodies = list(v8.element.body.iterchildren())
    equations = [v8_bodies[i] for i in V8_EQUATION_INDICES]
    for idx, el in zip(V8_EQUATION_INDICES, equations):
        assert el.find(M_NS + "oMath") is not None, f"body[{idx}] is not an equation"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 2.0      # double-spaced for review

    fig_caps: list[str] = []
    eq_no = 0
    table_no = 0
    pending_caption: str | None = None

    for tag, payload in blocks:
        if tag == "TITLE":
            add_par(doc, payload, size=14, bold=True)
        elif tag == "H1":
            add_par(doc, payload, style="Heading 1")
        elif tag == "H2":
            add_par(doc, payload, style="Heading 2")
        elif tag == "HU":
            add_par(doc, payload, style="Heading 1")
        elif tag == "P":
            add_par(doc, apply_citations(payload, order))
        elif tag == "EQ":
            el = copy.deepcopy(equations[eq_no])
            eq_no += 1
            doc.element.body.append(el)
            par = Paragraph(el, doc)
            par.runs[-1].text = f"\t({eq_no})"
        elif tag == "TABLECAP":
            table_no += 1
            pending_caption = apply_citations(payload, order)
            add_par(doc, pending_caption, bold=False, size=10)
        elif tag == "TABLE":
            rows = payload
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    cp = t.cell(r, c).paragraphs[0]
                    cp.paragraph_format.line_spacing = 1.0
                    cp.paragraph_format.space_after = Pt(0)
                    run = cp.add_run(apply_citations(cell, order))
                    run.font.size = Pt(9)
                    run.bold = (r == 0)
            doc.add_paragraph()
        elif tag == "FIGCAP":
            fig_caps.append(apply_citations(payload, order))
        else:
            raise SystemExit(f"unknown tag {tag}")

    # references, in citation order
    add_par(doc, "References", style="Heading 1")
    for key, num in sorted(order.items(), key=lambda kv: kv[1]):
        add_par(doc, f"[{num}] {refs[key]}", size=10)

    add_par(doc, "Figure captions", style="Heading 1")
    for cap in fig_caps:
        add_par(doc, cap, size=10)

    doc.save(DST)

    # ---- report
    words = 0
    for p in doc.paragraphs:
        words += len(p.text.split())
    tbl_words = sum(len(c.text.split())
                    for t in doc.tables for row in t.rows for c in row.cells)
    body_words = words - sum(
        len(p.text.split()) for p in doc.paragraphs
        if p.text.startswith("[") and re.match(r"^\[\d+\]", p.text))
    print(f"wrote {DST.name}")
    print(f"  sections numbered, {len(order)} references, {len(fig_caps)} figures, "
          f"{table_no} tables, {eq_no} equations")
    print(f"  words in paragraphs (incl. references): {words}")
    print(f"  words excluding the reference list     : {body_words}")
    print(f"  words inside tables                    : {tbl_words}")

    unused = sorted(set(refs) - set(order))
    print(f"  reference entries dropped from v8      : {len(unused)} -> {unused}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
